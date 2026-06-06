# -*- coding: utf-8 -*-
"""生成《心理港湾 · 系统详细说明》Word 文档（完整版）"""

import os
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DOCS_DIR = os.path.dirname(SCRIPT_DIR)
OUTPUT = os.path.join(DOCS_DIR, "心理港湾_系统详细说明.docx")

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    print("请先安装: pip install python-docx")
    sys.exit(1)


def set_cn_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def add_title(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_cn_font(run, "黑体", 16 if level == 1 else 14 if level == 2 else 12, bold=True)


def add_para(doc, text, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_cn_font(run)
    p.paragraph_format.line_spacing = 1.5


def add_bullet(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_cn_font(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        for run in p.runs:
            set_cn_font(run)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_cn_font(run, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()


def add_code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cn_font(run, "Consolas", 10)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.line_spacing = 1.25


def build():
    doc = Document()

    # 封面
    for _ in range(3):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn_font(t.add_run("心理港湾（PsyQA）"), "黑体", 28, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn_font(
        sub.add_run("\n高校双端智能心理健康服务平台\n\n系统详细说明文档"),
        "黑体",
        16,
        bold=True,
    )
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn_font(info.add_run("\n\n版本 1.2 · 2026年5月\n项目代号：xinli-gangwan"), "宋体", 12)

    doc.add_page_break()

    # 1 概述
    add_title(doc, "一、项目概述", 1)
    add_para(
        doc,
        "「心理港湾」是一套面向高校场景的校园心理健康智能咨询与院系管理平台。"
        "系统为学生提供 AI 辅助的心理倾诉与结构化评估，为辅导员提供院系态势看板与风险告警，"
        "为管理员提供全校报表导出与用户管理等能力。产品采用学生端与学校端双端分离架构，"
        "打通「咨询—识别—协同」的校园心理健康数字化原型。",
        indent=True,
    )
    add_para(doc, "设计目标包括：", indent=True)
    add_bullet(
        doc,
        [
            "温暖可用：对话式咨询、情绪识别、关怀建议与危机提示；",
            "数据可管：咨询记录、心理画像、趋势统计与告警闭环；",
            "权限分级：学生 / 辅导员 / 管理员分端访问，对话隐私可配置；",
            "智能可演示：默认优先本地大模型（Ollama）生成回复，报告与画像后台异步生成。",
        ],
    )

    # 2 功能
    add_title(doc, "二、系统功能", 1)

    add_title(doc, "2.1 学生端（/student）", 2)
    add_table(
        doc,
        ["功能模块", "说明"],
        [
            ["智能咨询", "多轮对话；知识库检索 + 规则心理分析 + 默认 Ollama 大模型生成回复"],
            ["大模型标识", "聊天区与报告区展示「大模型回复」标识（llmUsed）"],
            ["情绪与风险", "自动识别情绪、问题领域、风险等级；高危词触发安全话术与热线"],
            ["心理画像", "对话结束后异步生成；前端轮询自动刷新"],
            ["心理咨询报告", "趋势图、可折叠报告、情绪/风险/问题摘要；报告可异步补全"],
            ["报告导出", "复制摘要、分享辅导员（简报）、一页纸 PDF、导出完整 TXT"],
            ["个人进度", "历史摘要、分组历史、自评、趋势图"],
            ["个人资料", "昵称、头像、院系；可关闭「允许学校端查看对话原文」"],
            ["危机支持", "高危场景展示援助热线与支持横幅"],
        ],
    )

    add_title(doc, "2.2 学校端 · 辅导员（/school）", 2)
    add_table(
        doc,
        ["功能模块", "说明"],
        [
            ["数据看板", "累计咨询、活跃学生、高危人数、待处理告警、问题/情绪 Top3 等"],
            ["风险告警", "列表筛选、指派、备注、误报标记、状态流转"],
            ["学生列表", "本院系学生档案，默认脱敏展示（姓名/学号掩码）"],
            ["学生详情", "咨询摘要、心理指标时间线；对话原文受隐私开关控制"],
        ],
    )
    add_para(doc, "辅导员仅可查看其管辖院系（managedOrgIds）范围内的数据。", indent=True)

    add_title(doc, "2.3 学校端 · 管理员", 2)
    add_table(
        doc,
        ["功能模块", "说明"],
        [
            ["全校报表导出", "看板页导出 CSV / JSON（GET /api/school/admin/report）"],
            ["用户管理", "/school/admin/users：新建/编辑用户、角色、院系、重置密码"],
            ["知识库维护", "查看/触发知识库更新"],
            ["系统统计", "全局咨询统计"],
        ],
    )

    add_title(doc, "2.4 报告导出方式（内容互不重复）", 2)
    add_table(
        doc,
        ["操作", "内容说明"],
        [
            ["复制摘要", "结构化摘要 + 完整报告正文"],
            ["分享辅导员", "简报（情绪/风险/指标/关注建议/隐私说明），不含全文"],
            ["一页纸 PDF", "一页要点 + 核心指标 + 报告摘录（约 380 字），浏览器打印另存"],
            ["导出完整报告", "下载 .txt 全文"],
        ],
    )

    add_title(doc, "2.5 后台能力", 2)
    add_bullet(
        doc,
        [
            "知识库调度：支持定时更新、计划任务、管理员 API 触发；",
            "训练流水线：公开数据集拉取、语料构建、心理权重训练（Python 脚本）；",
            "演示数据种子：启动时注入 demo 学生咨询记录与告警样例。",
        ],
    )

    # 3 权限
    add_title(doc, "三、角色与权限", 1)
    add_table(
        doc,
        ["角色", "标识", "默认入口"],
        [
            ["学生", "student", "/student"],
            ["辅导员", "counselor", "/school/dashboard"],
            ["管理员", "admin", "/school/dashboard"],
        ],
    )
    add_para(doc, "主要权限标识（Permission）：", indent=True)
    add_table(
        doc,
        ["权限", "学生", "辅导员", "管理员"],
        [
            ["consult:own / report:own / trend:own", "✓", "", ""],
            ["school:dashboard / school:alerts:read", "", "✓", "✓"],
            ["school:students:masked", "", "✓", "✓"],
            ["school:transcripts:full", "", "可扩展", "✓（默认）"],
            ["school:manage", "", "", "✓"],
        ],
    )
    add_para(
        doc,
        "路由层通过 requireAuth、requireRole、requirePermission 中间件校验。"
        "管理员专属接口挂载在 /api/school/admin/*。"
        "学生可在资料中关闭 allowSchoolTranscriptView，学校端将不展示完整对话原文。",
        indent=True,
    )

    # 4 架构
    add_title(doc, "四、技术架构", 1)
    add_title(doc, "4.1 总体架构", 2)
    add_code_block(
        doc,
        "浏览器（React SPA：学生端 /student、学校端 /school、登录 /login）\n"
        "        │  HTTP REST · Cookie: psyqa_session · Bearer 兼容\n"
        "        ▼\n"
        "Express 后端（Node.js + TypeScript）\n"
        "  /api/auth   /api/questions   /api/school   /api/school/admin\n"
        "        │              │                │\n"
        "        ▼              ▼                ▼\n"
        "   SQLite/JSON    知识库+向量库      Ollama（默认启用）\n"
        "   账户/对话/告警  mental_dataset    qwen:7b / embed",
    )

    add_title(doc, "4.2 咨询回答流水线（核心）", 2)
    add_para(
        doc,
        "用户提问经 server/src/services/question/answerOrchestrator.ts 编排，"
        "采用「先回复、后报告」策略，突出大模型在对话生成中的作用：",
        indent=True,
    )
    add_numbered(
        doc,
        [
            "快速规则评估：emotionService 识别情绪、风险、问题类型；危机词（如自伤、打架等）命中则优先返回安全计划话术；",
            "知识检索：ragService 从 mental_dataset 检索参考条目，供大模型 prompt 使用；",
            "优先大模型回复：默认调用 Ollama（qwen:7b）生成共情与建议正文，聊天区展示「大模型回复」；",
            "即时返回：保存对话与初步心理快照（PsychSnapshot），返回占位报告；",
            "异步完整报告：reportQueue 后台执行 LLM 心理分析、统计融合（psychStatsService）、写入完整报告；",
            "异步咨询画像：portraitQueue 在报告就绪后生成会话画像（ConversationPortrait）。",
        ],
    )

    add_title(doc, "4.3 大模型相关模块", 2)
    add_table(
        doc,
        ["模块/文件", "职责"],
        [
            ["ollamaAvailability.ts", "控制是否启用大模型；默认开启，PSYQA_FAST_ANSWER=1 可关闭"],
            ["ollamaClient.ts", "Ollama HTTP 调用与健康检查"],
            ["answerOrchestrator.ts", "编排检索、大模型生成、保存与异步队列"],
            ["reportQueue.ts", "异步生成完整心理评估报告并更新数据库"],
            ["portraitQueue.ts", "异步生成咨询画像"],
            ["psychAnalysisService.ts", "可选 LLM 心理结构化分析（用于异步报告阶段）"],
        ],
    )

    add_title(doc, "4.4 数据存储", 2)
    add_table(
        doc,
        ["存储", "说明"],
        [
            ["SQLite（默认）", "server/data/psyqa.db；表：accounts、dialogs、alerts、meta 等"],
            ["JSON 回退", "设置 PSYQA_USE_JSON_STORAGE=1 使用 accounts.json、user_history.json"],
            ["迁移", "首次启用 SQLite 时自动从 JSON 迁移（migrateFromJson.ts）"],
        ],
    )

    add_title(doc, "4.5 认证与会话", 2)
    add_bullet(
        doc,
        [
            "登录：POST /api/auth/login，密码 scrypt 加盐哈希存储；",
            "会话：签名 Token，生产环境写入 HttpOnly Cookie psyqa_session；",
            "登出：POST /api/auth/logout 清除 Cookie；",
            "登录限流：loginRateLimit 防暴力尝试；",
            "生产环境必须配置 AUTH_SECRET（≥16 字符）。",
        ],
    )

    # 5 技术栈
    add_title(doc, "五、技术栈", 1)
    add_title(doc, "5.1 前端", 2)
    add_table(
        doc,
        ["技术", "说明"],
        [
            ["React 19", "学生端 / 学校端 SPA"],
            ["TypeScript", "类型安全"],
            ["React Router 7", "分路由"],
            ["Axios", "HTTP，withCredentials 支持 Cookie"],
            ["Create React App", "开发端口 3000，代理 3001"],
        ],
    )
    add_title(doc, "5.2 后端", 2)
    add_table(
        doc,
        ["技术", "说明"],
        [
            ["Node.js 18+", "运行环境"],
            ["Express 4", "REST API"],
            ["TypeScript 5", "服务端编译"],
            ["better-sqlite3", "嵌入式数据库"],
            ["Jest", "单元测试"],
        ],
    )
    add_title(doc, "5.3 AI 与算法", 2)
    add_table(
        doc,
        ["组件", "说明"],
        [
            ["Ollama", "本地大模型，默认 qwen:7b，用于对话生成与报告阶段心理分析"],
            ["nomic-embed-text", "默认嵌入模型，npm run build:embeddings 构建向量"],
            ["TF-IDF", "无嵌入时的向量检索兜底"],
            ["规则引擎", "情绪、风险、问题分类与高危词检测"],
            ["RAG", "mental_dataset.json 知识条目检索"],
            ["Python 脚本", "知识库扩充、训练语料、权重训练（scripts/）"],
        ],
    )

    # 6 目录
    add_title(doc, "六、主要目录结构", 1)
    add_code_block(
        doc,
        "PsyQA/\n"
        "├── client/src/          # React 前端（student/ school/ components/ api/）\n"
        "├── server/src/          # Express 后端\n"
        "│   ├── routes/          # auth、question、school、schoolAdmin\n"
        "│   ├── services/        # 业务逻辑\n"
        "│   │   ├── question/    # answerOrchestrator 等\n"
        "│   │   ├── reportQueue.ts\n"
        "│   │   └── portraitQueue.ts\n"
        "│   └── db/              # SQLite 与迁移\n"
        "├── server/data/         # 数据库、知识库、种子数据\n"
        "├── docs/                # 文档与本 Word 生成脚本\n"
        "└── scripts/             # Python 训练与知识库脚本",
    )

    # 7 API
    add_title(doc, "七、API 概览", 1)
    add_title(doc, "7.1 认证 /api/auth", 2)
    add_table(
        doc,
        ["方法", "路径", "说明"],
        [
            ["POST", "/register", "学生注册"],
            ["POST", "/login", "登录"],
            ["POST", "/logout", "登出"],
            ["GET", "/me", "当前用户"],
            ["PATCH", "/profile", "学生资料（含隐私开关）"],
        ],
    )
    add_title(doc, "7.2 咨询 /api/questions", 2)
    add_table(
        doc,
        ["方法", "路径", "说明"],
        [
            ["POST", "/ask", "发起咨询（返回 answer、report、portraitPending、reportPending）"],
            ["POST", "/analyze", "心理分析"],
            ["GET", "/progress", "个人进度与 latestAssessment"],
            ["GET", "/history/grouped", "分组历史"],
        ],
    )
    add_title(doc, "7.3 学校端 /api/school 与管理员 /api/school/admin", 2)
    add_table(
        doc,
        ["方法", "路径", "说明"],
        [
            ["GET", "/dashboard", "看板统计"],
            ["GET", "/students", "学生列表"],
            ["GET", "/alerts", "告警列表"],
            ["GET", "/admin/report", "全校报表 CSV/JSON"],
            ["GET/POST/PATCH", "/admin/users", "用户管理"],
        ],
    )

    # 8 演示账号
    add_title(doc, "八、演示账号", 1)
    add_table(
        doc,
        ["角色", "用户名", "密码", "说明"],
        [
            ["学生", "demo", "demo123", "含预置咨询记录"],
            ["辅导员", "counselor", "counselor123", "院系看板与告警"],
            ["管理员", "admin", "admin123", "报表导出与用户管理"],
            ["批量学生", "student01～20", "Student123456", "压测/演示账号"],
        ],
    )

    # 9 安装运行
    add_title(doc, "九、安装与运行", 1)
    add_title(doc, "9.1 环境要求", 2)
    add_bullet(
        doc,
        [
            "Node.js >= 18（建议 20+）；",
            "Ollama（推荐）：用于大模型对话与可选嵌入；",
            "Python 3.8+（可选）：生成 Word 文档或运行训练脚本。",
        ],
    )
    add_title(doc, "9.2 快速启动", 2)
    add_code_block(
        doc,
        "cd PsyQA\n"
        "npm run install:all\n"
        "npm run dev\n"
        "\n"
        "前端：http://localhost:3000\n"
        "后端：http://localhost:3001",
    )
    add_title(doc, "9.3 Ollama 配置（推荐）", 2)
    add_code_block(
        doc,
        "ollama pull qwen:7b\n"
        "ollama serve",
    )
    add_para(
        doc,
        "系统默认优先使用大模型生成咨询回复。若仅需快速演示、不调用 Ollama，"
        "可在 .env 中设置 PSYQA_FAST_ANSWER=1 或 PSYQA_SKIP_OLLAMA=1 后重启后端。",
        indent=True,
    )

    add_title(doc, "9.4 常用环境变量", 2)
    add_table(
        doc,
        ["变量", "说明"],
        [
            ["AUTH_SECRET", "会话密钥（生产必填，≥16 字符）"],
            ["PORT", "后端端口，默认 3001"],
            ["CORS_ORIGIN", "前端域名（Cookie 跨域时需配置）"],
            ["NODE_ENV", "production 时启用安全策略与静态托管"],
            ["PSYQA_FAST_ANSWER=1", "关闭大模型，使用规则+知识库快速回复"],
            ["PSYQA_SKIP_OLLAMA=1", "跳过 Ollama 调用"],
            ["PSYQA_USE_JSON_STORAGE=1", "使用 JSON 而非 SQLite"],
            ["OLLAMA_API_URL", "Ollama 地址，默认 http://localhost:11434"],
            ["OLLAMA_MODEL", "对话模型，默认 qwen:7b"],
            ["OLLAMA_EMBED_MODEL", "嵌入模型，默认 nomic-embed-text"],
        ],
    )

    add_title(doc, "9.5 生成 Word 文档", 2)
    add_code_block(
        doc,
        "pip install python-docx\n"
        "python docs/scripts/generate_system_doc.py\n"
        "python docs/scripts/generate_work_doc.py",
    )
    add_para(
        doc,
        "将分别生成：docs/心理港湾_系统详细说明.docx（本文档）、"
        "docs/心理港湾_大学生作品说明.docx（精简答辩版）。",
        indent=True,
    )

    # 10 安全
    add_title(doc, "十、安全说明", 1)
    add_bullet(
        doc,
        [
            "密码使用 scrypt 加盐哈希，不明文存储；",
            "生产环境强制 AUTH_SECRET；Cookie 在生产为 Secure；",
            "响应头：X-Content-Type-Options、X-Frame-Options、Referrer-Policy；",
            "学生咨询接口 requireStudentAccess：仅本人或开发游客可访问；",
            "学校端需登录且角色为辅导员或管理员；管理接口额外校验 school:manage；",
            "登录接口限流，降低撞库风险。",
        ],
    )

    # 11 伦理
    add_title(doc, "十一、伦理与使用声明", 1)
    add_para(
        doc,
        "本系统为教学展示与校园心理健康信息化原型，不能替代专业心理咨询或医疗诊断。"
        "涉及自伤、自杀、暴力冲突等危机内容时，请立即联系校心理中心或拨打心理援助热线"
        "（如全国心理援助热线 400-161-9995）。"
        "学校端查看学生信息须遵守法律法规及学校隐私制度；"
        "学生可在个人资料中控制对话原文是否对学校可见。",
        indent=True,
    )

    # 12 文档索引
    add_title(doc, "十二、相关文档索引", 1)
    add_table(
        doc,
        ["文档", "路径", "内容"],
        [
            ["系统说明（Markdown）", "docs/系统说明文档.md", "与本文档同步的在线版"],
            ["系统详细说明（Word）", "docs/心理港湾_系统详细说明.docx", "本文档"],
            ["作品说明（Word）", "docs/心理港湾_大学生作品说明.docx", "精简答辩版"],
            ["安装与运行", "docs/安装与运行.md", "环境、Ollama、端口"],
            ["演示说明", "docs/作品使用与演示说明.md", "演示流程"],
            ["答辩脚本", "docs/03_演示汇报脚本.md", "口播脚本"],
            ["架构图", "docs/系统架构图.txt", "文本架构图"],
        ],
    )

    add_title(doc, "十三、版本信息", 1)
    add_table(
        doc,
        ["项目", "说明"],
        [
            ["产品名称", "心理港湾"],
            ["npm 包名", "xinli-gangwan"],
            ["文档版本", "1.2（2026-05）"],
            ["核心模块", "answerOrchestrator、reportQueue、portraitQueue、ReportPanel"],
        ],
    )
    add_para(
        doc,
        "文档随代码演进更新；若与实现不一致，以 server/src 与 client/src 源码为准。",
        indent=True,
    )

    doc.save(OUTPUT)
    print(f"已生成: {OUTPUT}")


if __name__ == "__main__":
    build()
