# -*- coding: utf-8 -*-
"""生成《心理港湾 · 大学生作品说明》Word 文档"""

import os
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DOCS_DIR = os.path.dirname(SCRIPT_DIR)
OUTPUT = os.path.join(DOCS_DIR, "心理港湾_大学生作品说明.docx")

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    print("请先安装: pip install python-docx")
    sys.exit(1)


def set_cn_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def add_title(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_cn_font(run, "黑体", 16 if level == 1 else 14 if level == 2 else 12, bold=True)


def add_para(doc, text, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_cn_font(run)
    p.paragraph_format.line_spacing = 1.5


def add_bullet(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_cn_font(run)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()


def build():
    doc = Document()

    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn_font(t.add_run("心理港湾"), "黑体", 26, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn_font(sub.add_run("\n高校双端智能心理健康服务平台\n\n大学生作品说明"), "黑体", 15, bold=True)
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn_font(info.add_run("\n\n作品版本 V2.1 · 2026年5月"), "宋体", 12)

    doc.add_page_break()

    add_title(doc, "一、作品概述", 1)
    add_para(
        doc,
        "「心理港湾」是本团队面向高校场景设计的大学生作品，适用于毕业设计、创新训练、课程设计等。"
        "系统采用学生端与学校端双端分离架构，打通「咨询—识别—协同」的校园心理健康数字化原型。"
        "学生可获得 AI 辅助倾诉与结构化心理报告；学校端在合规前提下掌握院系风险态势与告警。",
        indent=True,
    )

    add_title(doc, "二、主要功能", 1)

    add_title(doc, "2.1 学生端", 2)
    add_bullet(doc, [
        "智能咨询：默认 Ollama 大模型生成回复；知识库检索 + 规则分析；报告与画像后台生成",
        "心理咨询报告：趋势图、会话画像、可折叠报告与情绪/风险/问题摘要",
        "报告获取：复制摘要、分享辅导员（简报）、一页纸 PDF、导出完整报告",
        "隐私设置：可关闭「允许学校端查看对话原文」",
        "危机提示：高危场景展示援助热线与支持信息",
    ])

    add_title(doc, "2.2 学校端", 2)
    add_bullet(doc, [
        "辅导员：数据看板、风险告警、脱敏学生列表与档案",
        "管理员：全校报表导出（CSV/JSON）、用户管理（新建/编辑/重置密码）",
        "权限隔离：辅导员仅管辖院系；管理接口需 admin + school:manage",
    ])

    add_title(doc, "三、技术架构", 1)
    add_bullet(doc, [
        "前端：React 19 + TypeScript + React Router",
        "后端：Node.js + Express + TypeScript",
        "数据：SQLite 默认（server/data/psyqa.db），可回退 JSON",
        "AI：Ollama（qwen:7b）默认优先；RAG 知识库 + 规则引擎；reportQueue 异步报告",
        "安全：scrypt 密码哈希、HttpOnly Cookie 会话、登录限流、角色权限中间件",
    ])

    add_title(doc, "四、演示账号", 1)
    add_table(doc, ["角色", "用户名", "密码", "说明"], [
        ["学生", "demo", "demo123", "含预置咨询与告警样例"],
        ["辅导员", "counselor", "counselor123", "院系看板与告警"],
        ["管理员", "admin", "admin123", "报表导出与用户管理"],
    ])

    add_title(doc, "五、运行方式", 1)
    add_bullet(doc, [
        "进入目录 PsyQA，执行：npm run install:all",
        "启动：npm run dev",
        "浏览器访问：http://localhost:3000/login",
        "推荐：安装 Ollama 并 ollama pull qwen:7b（默认大模型回复）",
        "快速模式：.env 设置 PSYQA_FAST_ANSWER=1",
        "详细说明见 docs/心理港湾_系统详细说明.docx、docs/系统说明文档.md",
    ])

    add_title(doc, "六、推荐演示顺序", 1)
    add_bullet(doc, [
        "学生 demo 登录 → 伦理确认 → 咨询 → 展开报告与四种导出方式",
        "辅导员 counselor 登录 → 看板 → 告警 → 学生列表",
        "管理员 admin 登录 → 导出全校报表 → 用户管理",
    ])

    add_title(doc, "七、伦理声明", 1)
    add_para(
        doc,
        "本系统为教学展示原型，不能替代专业心理咨询或医疗诊断。"
        "涉及自伤、自杀等危机内容时，请立即联系校心理中心或拨打心理援助热线。"
        "学校端查看学生信息须遵守法律法规及学校隐私制度；学生可在个人资料中控制对话原文是否对学校可见。",
        indent=True,
    )

    add_title(doc, "八、文档索引", 1)
    add_table(doc, ["文档", "路径"], [
        ["系统详细说明（Word）", "docs/心理港湾_系统详细说明.docx"],
        ["系统说明（Markdown）", "docs/系统说明文档.md"],
        ["演示说明", "docs/作品使用与演示说明.md"],
        ["答辩脚本", "docs/03_演示汇报脚本.md"],
        ["安装运行", "docs/安装与运行.md"],
        ["架构图", "docs/系统架构图.txt"],
    ])

    doc.save(OUTPUT)
    print(f"已生成: {OUTPUT}")


if __name__ == "__main__":
    build()
