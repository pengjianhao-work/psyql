#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 training_data/ 下各开源数据集归一化为统一训练语料

输出:
  training_data/unified_corpus.jsonl   — 全量（含伪标签，供算法训练）
  training_data/unified_corpus_stats.json
  training_data/knowledge_candidates.json — 供 expand_knowledge_base 合并

用法:
  python scripts/build_training_corpus.py
  python scripts/build_training_corpus.py --max-records 50000
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set

from psych_corpus_utils import (
    build_record,
    clean_text,
    efaqa_label_scores,
    infer_cognitive_distortions,
    infer_risk_level,
)

try:
    from dataset_network import link_psyqa_from_repo, scan_local_inventory
except ImportError:
    link_psyqa_from_repo = lambda: False  # type: ignore
    scan_local_inventory = lambda: {"ready": [], "missing": []}  # type: ignore

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
DATA_ROOT = REPO_ROOT / "training_data"
OUT_JSONL = DATA_ROOT / "unified_corpus.jsonl"
OUT_STATS = DATA_ROOT / "unified_corpus_stats.json"
OUT_KNOWLEDGE = DATA_ROOT / "knowledge_candidates.json"
REGISTRY = SCRIPT_DIR / "dataset_registry.json"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def load_jsonl(path: Path) -> List[Dict[str, Any]]:
    rows = []
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def parse_psyqa(path: Path) -> Iterable[Dict[str, Any]]:
    for item in load_json(path):
        q = item.get("question", "")
        answers = item.get("answers") or []
        if not q or not answers:
            continue
        a = answers[0].get("answer_text", "") if isinstance(answers[0], dict) else str(answers[0])
        rec = build_record(
            "psyqa",
            q,
            a,
            algorithm_targets=["statistical", "ml", "neural"],
        )
        if rec:
            yield rec


def parse_efaqa(path: Path) -> Iterable[Dict[str, Any]]:
    data = load_json(path)
    for item in data:
        title = item.get("title", "")
        desc = item.get("description", "")
        label = item.get("label") or {}
        chats = item.get("chats") or []
        turns = []
        for c in chats:
            sender = c.get("sender", "")
            role = "user" if sender == "owner" else "counselor"
            val = c.get("value", "")
            if val:
                turns.append({"role": role, "text": clean_text(val, 500)})
        user_msgs = [t["text"] for t in turns if t["role"] == "user"]
        counselor_msgs = [t["text"] for t in turns if t["role"] == "counselor"]
        q = clean_text(f"{title} {desc}".strip() or (user_msgs[0] if user_msgs else ""), 400)
        a = clean_text(" ".join(counselor_msgs[-3:]) if counselor_msgs else "", 900)
        if not a and counselor_msgs:
            a = counselor_msgs[-1]
        scores = efaqa_label_scores(label)
        rec = build_record(
            "efaqa",
            q,
            a,
            turns=turns[:40],
            labels=scores,
            algorithm_targets=["statistical", "ml", "neural", "risk"],
        )
        if rec:
            yield rec


def parse_soulchat_dir(root: Path) -> Iterable[Dict[str, Any]]:
    for fp in root.rglob("*.json"):
        if fp.stat().st_size > 200_000_000:
            continue
        try:
            data = load_json(fp)
        except Exception:
            continue
        items = data if isinstance(data, list) else data.get("data", data.get("samples", []))
        if not isinstance(items, list):
            if isinstance(data, dict) and ("instruction" in data or "input" in data):
                items = [data]
            else:
                continue
        for item in items:
            text = item.get("instruction") or item.get("input") or item.get("text") or ""
            output = item.get("output") or item.get("response") or ""
            if not output and "conversations" in item:
                conv = item["conversations"]
                user_parts = []
                bot_parts = []
                for turn in conv:
                    role = turn.get("role", turn.get("from", ""))
                    content = turn.get("content", turn.get("value", ""))
                    if role in ("user", "human"):
                        user_parts.append(content)
                    else:
                        bot_parts.append(content)
                text = user_parts[0] if user_parts else ""
                output = bot_parts[-1] if bot_parts else ""
            if "用户：" in text or "心理咨询师：" in text:
                parts = re.split(r"用户：|心理咨询师：", text)
                user_chunks = [p.strip() for i, p in enumerate(parts) if i % 2 == 1 or (i > 0 and i % 2 == 0)]
                if len(parts) > 2:
                    q = parts[1].strip() if len(parts) > 1 else text[:200]
                    a = output or (parts[-1].strip() if parts else "")
                else:
                    q = text[:300]
                    a = output
            else:
                q = text[:300]
                a = output
            rec = build_record(
                "soulchat",
                q,
                a,
                algorithm_targets=["ml", "neural", "timeseries"],
            )
            if rec:
                yield rec


def parse_mentalchat(path: Path) -> Iterable[Dict[str, Any]]:
    rows = load_jsonl(path) if path.suffix == ".jsonl" else []
    if not rows and path.suffix == ".json":
        raw = load_json(path)
        rows = raw if isinstance(raw, list) else []
    for row in rows:
        q = row.get("question") or row.get("input") or row.get("patient") or row.get("instruction") or ""
        a = row.get("answer") or row.get("output") or row.get("counselor") or row.get("response") or ""
        if isinstance(q, list):
            q = " ".join(str(x) for x in q)
        if isinstance(a, list):
            a = " ".join(str(x) for x in a)
        topic = str(row.get("topic", row.get("category", "")))
        labels = {}
        if topic:
            labels["problemTypes"] = [topic[:40]]
        risk = infer_risk_level(f"{q} {a}")
        labels["riskLevel"] = risk
        rec = build_record(
            "mentalchat16k",
            str(q),
            str(a),
            labels=labels,
            algorithm_targets=["statistical", "ml", "neural"],
            language="en",
        )
        if rec:
            yield rec


def parse_psydial(path: Path) -> Iterable[Dict[str, Any]]:
    data = load_json(path)
    dialogues = data if isinstance(data, list) else data.get("data", data.get("dialogues", []))
    for dlg in dialogues:
        turns_raw = dlg if isinstance(dlg, list) else dlg.get("dialogue", dlg.get("turns", dlg.get("conversation", [])))
        if isinstance(dlg, dict) and not turns_raw:
            turns_raw = dlg.get("messages", [])
        turns: List[Dict[str, str]] = []
        for t in turns_raw or []:
            if isinstance(t, dict):
                role = t.get("role", t.get("speaker", "user"))
                text = t.get("content", t.get("text", t.get("utterance", "")))
            else:
                role, text = "user", str(t)
            r = "counselor" if str(role).lower() in ("counselor", "assistant", "therapist", "咨询师") else "user"
            turns.append({"role": r, "text": clean_text(str(text), 500)})
        user_msgs = [x["text"] for x in turns if x["role"] == "user"]
        bot_msgs = [x["text"] for x in turns if x["role"] == "counselor"]
        if not user_msgs:
            continue
        q = user_msgs[0]
        a = bot_msgs[-1] if bot_msgs else ""
        rec = build_record(
            "psydial",
            q,
            a,
            turns=turns,
            algorithm_targets=["timeseries", "neural"],
        )
        if rec:
            yield rec


def parse_psycrisis_xlsx(path: Path) -> Iterable[Dict[str, Any]]:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        print("[WARN] psycrisis 需 openpyxl: pip install openpyxl")
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    text_idx = next(
        (i for i, h in enumerate(headers) if h in ("text", "prompt", "question", "user", "content")),
        0,
    )
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[text_idx]:
            continue
        text = str(row[text_idx])
        rec = build_record(
            "psycrisis",
            text[:300],
            "若你正经历危机，请立即联系学校心理中心、心理援助热线或身边可信的人。",
            labels={"riskLevel": "critical", "stressScore": 95, "anxietyScore": 92, "moodScore": 15},
            algorithm_targets=["risk", "ml", "neural"],
            language="en",
        )
        if rec:
            yield rec


def parse_hf_jsonl(path: Path, source: str, default_targets: List[str], language: str = "zh") -> Iterable[Dict[str, Any]]:
    rows = load_jsonl(path) if path.suffix == ".jsonl" else []
    if not rows and path.suffix == ".json":
        raw = load_json(path)
        rows = raw if isinstance(raw, list) else [raw]
    for row in rows:
        q = (
            row.get("question")
            or row.get("input")
            or row.get("text")
            or row.get("content")
            or row.get("instruction")
            or row.get("query")
            or row.get("post")
            or ""
        )
        a = (
            row.get("answer")
            or row.get("output")
            or row.get("response")
            or row.get("reply")
            or row.get("label_text")
            or ""
        )
        if not a and row.get("reasoning"):
            a = str(row.get("reasoning", ""))[:600]
        if isinstance(q, list):
            q = " ".join(str(x) for x in q)
        if isinstance(a, list):
            a = " ".join(str(x) for x in a)
        labels: Dict[str, Any] = {}
        risk_label = row.get("risk") or row.get("risk_level") or row.get("suicide_risk") or row.get("label")
        if risk_label is not None:
            rl = str(risk_label).lower()
            if rl in ("1", "high", "true", "yes", "高危", "高"):
                labels["riskLevel"] = "critical" if "自杀" in str(q) else "high"
            else:
                labels["riskLevel"] = infer_risk_level(str(q))
        if row.get("stress") is not None:
            labels["stressScore"] = float(row["stress"]) * (100 if float(row["stress"]) <= 1 else 1)
        if row.get("distortion") or row.get("cognitive_distortion"):
            labels["cognitiveDistortions"] = [str(row.get("distortion") or row.get("cognitive_distortion"))]
        if not labels.get("cognitiveDistortions"):
            labels["cognitiveDistortions"] = infer_cognitive_distortions(f"{q} {a}")
        rec = build_record(source, str(q), str(a), labels=labels or None, algorithm_targets=default_targets, language=language)
        if rec:
            yield rec


def parse_sos_hl(path: Path) -> Iterable[Dict[str, Any]]:
    for row in load_jsonl(path) if path.suffix == ".jsonl" else []:
        text = str(row.get("text") or row.get("content") or row.get("post") or "")
        label = row.get("label") if row.get("label") is not None else row.get("suicide_risk") or row.get("risk")
        if str(label) in ("1", 1, "high", "High", "高危"):
            risk = "critical"
        elif str(label) in ("0", 0, "low", "Low", "低危"):
            risk = "low"
        else:
            risk = "medium"
        rec = build_record(
            "sos_hl_1k",
            text[:400],
            "你并不孤单。请尽快联系学校心理中心、12356 心理援助热线或身边可信的人。",
            labels={"riskLevel": risk, "stressScore": 90 if risk == "critical" else 65, "anxietyScore": 88, "moodScore": 20},
            algorithm_targets=["risk", "ml"],
        )
        if rec:
            yield rec
    if path.suffix == ".jsonl":
        return
    yield from parse_hf_jsonl(path, "sos_hl_1k", ["risk", "ml"])


def parse_qiaoban(path: Path) -> Iterable[Dict[str, Any]]:
    for row in load_jsonl(path):
        q = row.get("question", "")
        a = row.get("answer", "")
        if not q or not a:
            continue
        rec = build_record(
            row.get("source", "qiaoban"),
            q,
            a,
            turns=row.get("turns"),
            labels=row.get("labels"),
            algorithm_targets=row.get("algorithmTargets"),
            language=row.get("language", "zh"),
        )
        if rec:
            if row.get("id"):
                rec["id"] = row["id"]
            yield rec


def parse_socialcd(path: Path) -> Iterable[Dict[str, Any]]:
    rows = load_jsonl(path) if path.suffix == ".jsonl" else []
    if not rows:
        yield from parse_hf_jsonl(path, "socialcd_3k", ["cognitive", "ml", "neural"])
        return
    for row in rows:
        text = str(row.get("text") or row.get("content") or row.get("sentence") or "")
        distortions = row.get("labels") or row.get("distortion") or row.get("cognitive_distortion")
        if isinstance(distortions, list):
            cog = [str(x) for x in distortions[:5]]
        elif distortions:
            cog = [str(distortions)]
        else:
            cog = infer_cognitive_distortions(text)
        rec = build_record(
            "socialcd_3k",
            text[:400],
            "我注意到你的想法里可能有一些常见的认知模式。我们可以一起慢慢梳理，不必苛责自己。",
            labels={"cognitiveDistortions": cog, "riskLevel": "low", "anxietyScore": 55},
            algorithm_targets=["cognitive", "ml", "neural"],
        )
        if rec:
            yield rec


def parse_modelscope_dir(root: Path, source: str, targets: List[str]) -> Iterable[Dict[str, Any]]:
    for fp in sorted(root.rglob("*.json*")):
        if fp.stat().st_size > 300_000_000:
            continue
        if fp.suffix == ".jsonl":
            yield from parse_hf_jsonl(fp, source, targets)
        elif fp.suffix == ".json":
            try:
                data = load_json(fp)
            except Exception:
                continue
            if isinstance(data, list) and data and isinstance(data[0], dict):
                if "instruction" in data[0] or "input" in data[0]:
                    for item in data:
                        q = item.get("instruction") or item.get("input") or ""
                        a = item.get("output") or item.get("response") or ""
                        raw_a = str(a)
                        think_end = "<" + "/think>"
                        if think_end in raw_a:
                            a = raw_a.split(think_end)[-1].strip()
                        rec = build_record(source, str(q), str(a), algorithm_targets=targets)
                        if rec:
                            yield rec
                else:
                    for item in data[:5000]:
                        yield from parse_hf_jsonl_from_dict(item, source, targets)


def parse_hf_jsonl_from_dict(row: Dict[str, Any], source: str, targets: List[str]) -> Iterable[Dict[str, Any]]:
    q = str(row.get("input") or row.get("question") or "")
    a = str(row.get("output") or row.get("answer") or "")
    rec = build_record(source, q, a, algorithm_targets=targets)
    if rec:
        yield rec


def collect_all_sources() -> tuple[List[Dict[str, Any]], Dict[str, int]]:
    all_records: List[Dict[str, Any]] = []
    sources_stats: Dict[str, int] = {}

    psyqa_path = DATA_ROOT / "psyqa" / "PsyQA_full.json"
    if not psyqa_path.exists():
        psyqa_path = REPO_ROOT / "PsyQA" / "PsyQA_full.json"
    if psyqa_path.exists():
        batch = list(parse_psyqa(psyqa_path))
        sources_stats["psyqa"] = len(batch)
        all_records.extend(batch)

    efaqa_path = DATA_ROOT / "efaqa" / "efaqa_corpus.json"
    if efaqa_path.exists():
        batch = list(parse_efaqa(efaqa_path))
        sources_stats["efaqa"] = len(batch)
        all_records.extend(batch)

    soul_dir = DATA_ROOT / "soulchat"
    if soul_dir.exists():
        batch = list(parse_soulchat_dir(soul_dir))
        sources_stats["soulchat"] = len(batch)
        all_records.extend(batch)

    mc_path = DATA_ROOT / "mentalchat16k" / "samples.jsonl"
    if mc_path.exists():
        batch = list(parse_mentalchat(mc_path))
        sources_stats["mentalchat16k"] = len(batch)
        all_records.extend(batch)

    pd_path = DATA_ROOT / "psydial" / "PsyDial-D3.json"
    if pd_path.exists():
        batch = list(parse_psydial(pd_path))
        sources_stats["psydial"] = len(batch)
        all_records.extend(batch)

    pc_path = DATA_ROOT / "psycrisis" / "PsyCrisis-Bench.xlsx"
    if pc_path.exists():
        batch = list(parse_psycrisis_xlsx(pc_path))
        sources_stats["psycrisis"] = len(batch)
        all_records.extend(batch)

    for sid, targets, lang in [
        ("sos_hl_1k", ["risk", "ml"], "zh"),
        ("socialcd_3k", ["cognitive", "ml", "neural"], "zh"),
    ]:
        samples = DATA_ROOT / sid / "samples.jsonl"
        if samples.exists():
            batch = list(parse_sos_hl(samples) if sid == "sos_hl_1k" else parse_socialcd(samples))
            sources_stats[sid] = len(batch)
            all_records.extend(batch)

    for sid, targets in [
        ("mental_r1_10k", ["ml", "neural", "statistical"]),
        ("tcci_lingxi", ["statistical", "ml", "neural"]),
        ("mental_9k", ["statistical", "ml", "neural"]),
    ]:
        root = DATA_ROOT / sid
        if root.exists():
            batch = list(parse_modelscope_dir(root, sid, targets))
            sources_stats[sid] = len(batch)
            all_records.extend(batch)

    qiaoban_samples = DATA_ROOT / "qiaoban" / "samples.jsonl"
    if qiaoban_samples.exists():
        batch = list(parse_qiaoban(qiaoban_samples))
        sources_stats["qiaoban"] = len(batch)
        all_records.extend(batch)

    dreaddir = DATA_ROOT / "dreaddit"
    if dreaddir.exists():
        n = 0
        for fp in dreaddir.rglob("*.csv"):
            try:
                import csv

                with open(fp, encoding="utf-8", errors="replace") as cf:
                    for row in csv.DictReader(cf):
                        text = str(row.get("text") or row.get("post") or "")
                        stress = row.get("label") or row.get("stress")
                        labels = {"stressScore": 75.0 if str(stress) in ("1", "True", "true") else 40.0}
                        rec = build_record(
                            "dreaddit",
                            text[:300],
                            "压力感受很常见，可以尝试规律作息、适度运动，并与信任的人倾诉。",
                            labels=labels,
                            language="en",
                            algorithm_targets=["ml", "statistical"],
                        )
                        if rec:
                            all_records.append(rec)
                            n += 1
            except Exception:
                pass
        if n:
            sources_stats["dreaddit"] = n

    return all_records, sources_stats


def dedupe_records(records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    seen: Set[str] = set()
    out: List[Dict[str, Any]] = []
    for r in records:
        key = re.sub(r"\s+", "", r["question"].lower())[:80]
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--max-records", type=int, default=80000)
    args = parser.parse_args()

    link_psyqa_from_repo()
    inv = scan_local_inventory()
    print(f"本地数据就绪: {', '.join(inv.get('ready', [])) or '无（将尝试仅用已有文件）'}")

    all_records, sources_stats = collect_all_sources()
    all_records = dedupe_records(all_records)
    if len(all_records) > args.max_records:
        all_records = all_records[: args.max_records]

    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    with open(OUT_JSONL, "w", encoding="utf-8") as f:
        for r in all_records:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    knowledge = [
        {"question": r["question"], "answer": r["answer"], "source": r["source"]}
        for r in all_records
        if r.get("language", "zh") == "zh"
    ][: min(8000, len(all_records))]
    OUT_KNOWLEDGE.write_text(json.dumps({"knowledge": knowledge}, ensure_ascii=False, indent=2), encoding="utf-8")

    risk_dist: Dict[str, int] = {}
    for r in all_records:
        lvl = r.get("labels", {}).get("riskLevel", "low")
        risk_dist[lvl] = risk_dist.get(lvl, 0) + 1

    stats = {
        "total": len(all_records),
        "bySource": sources_stats,
        "riskDistribution": risk_dist,
        "outputs": {
            "corpus": str(OUT_JSONL),
            "knowledgeCandidates": str(OUT_KNOWLEDGE),
        },
    }
    OUT_STATS.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"统一语料: {len(all_records)} 条")
    print(f"  来源: {sources_stats}")
    print(f"  风险分布: {risk_dist}")
    print(f"  -> {OUT_JSONL}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
